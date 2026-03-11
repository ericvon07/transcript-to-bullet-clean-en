"""
Read input files and write output files.
"""

import os
import openpyxl
from datetime import datetime


def read_transcripts(filepath):
    """Read the input xlsx file and return a list of company dicts with keys: company, positive, negative, dealbreaker."""
    if not os.path.exists(filepath):
        print(f"Error: input file not found at '{filepath}'")
        return []

    workbook = openpyxl.load_workbook(filepath)
    sheet = workbook.active

    headers = [cell.value for cell in sheet[1]]

    column_map = {
        "company": "company",
        "positive_points": "positive",
        "negative_points": "negative",
        "dealbreakers": "dealbreaker",
    }

    key_to_index = {}
    for xlsx_col, our_key in column_map.items():
        if xlsx_col in headers:
            key_to_index[our_key] = headers.index(xlsx_col)

    transcripts = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        entry = {}
        for our_key, index in key_to_index.items():
            entry[our_key] = row[index] if row[index] is not None else ""
        transcripts.append(entry)

    return transcripts


def write_output(results, output_dir):
    """Save the pipeline results to a timestamped xlsx file inside output_dir."""
    if not os.path.exists(output_dir):
        os.makedirs(output_dir)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"output_{timestamp}.xlsx"
    output_path = os.path.join(output_dir, filename)

    workbook = openpyxl.Workbook()
    sheet = workbook.active

    sheet.append(["company", "bullet_point", "macro_topic", "sentiment"])

    for item in results:
        sheet.append(
            [
                item["company"],
                item["bullet_point"],
                item["macro_topic"],
                item["sentiment"],
            ]
        )

    workbook.save(output_path)
    print(f"Output saved to: {output_path}")

    return output_path
